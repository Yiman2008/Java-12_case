# -*- coding: utf-8 -*-
"""Small python-docx wrapper with Russian academic (ГОСТ-like) formatting."""
import subprocess, re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SERIF = "Times New Roman"
MONO = "Courier New"


def set_font(run_or_style, name, size=None, bold=None, italic=None):
    f = run_or_style.font
    f.name = name
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    el = run_or_style.element if hasattr(run_or_style, "element") else run_or_style._element
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), name)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color)
    tcPr.append(sh)


def para_shade(p, color):
    pPr = p._p.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color)
    pPr.append(sh)


class Doc:
    def __init__(self):
        self.d = Document()
        self.fig = 0
        self.tab = 0
        self.headings = []  # (level, text)
        self.toc_pages = {}
        self._toc_par = None
        s = self.d.sections[0]
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.left_margin, s.right_margin = Cm(3.0), Cm(1.5)
        s.top_margin, s.bottom_margin = Cm(2.0), Cm(2.0)
        s.different_first_page_header_footer = True
        self._footer(s)
        st = self.d.styles
        n = st["Normal"]
        set_font(n, SERIF, 14)
        pf = n.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(0); pf.space_before = Pt(0)
        pf.first_line_indent = Cm(1.25)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for lvl, size in ((1, 16), (2, 14), (3, 14)):
            h = st[f"Heading {lvl}"]
            set_font(h, SERIF, size, bold=True, italic=False)
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.paragraph_format.first_line_indent = Cm(1.25) if lvl > 1 else Cm(0)
            h.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.line_spacing = 1.15
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if lvl == 1 else WD_ALIGN_PARAGRAPH.LEFT
        lb = st["List Bullet"]
        set_font(lb, SERIF, 14)
        lb.paragraph_format.first_line_indent = Cm(0)
        lb.paragraph_format.line_spacing = 1.5
        ln = st["List Number"]
        set_font(ln, SERIF, 14)

    # ------------------------------------------------------------ helpers
    def _footer(self, section):
        f = section.footer
        p = f.paragraphs[0] if f.paragraphs else f.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run()
        set_font(r, SERIF, 12)
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); r._r.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)

    def p(self, text="", bold=False, italic=False, align=None, indent=True, size=None, space_after=None, keep=False):
        par = self.d.add_paragraph()
        if not indent: par.paragraph_format.first_line_indent = Cm(0)
        if align == "center": par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right": par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "left": par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if space_after is not None: par.paragraph_format.space_after = Pt(space_after)
        if keep: par.paragraph_format.keep_with_next = True
        self._runs(par, text, bold, italic, size)
        return par

    def _runs(self, par, text, bold=False, italic=False, size=None):
        # **bold** inline markup
        parts = re.split(r"(\*\*.+?\*\*)", text)
        for part in parts:
            if not part: continue
            b = bold
            if part.startswith("**") and part.endswith("**"):
                part, b = part[2:-2], True
            r = par.add_run(part)
            set_font(r, SERIF, size, bold=b, italic=italic)

    def h(self, level, text):
        par = self.d.add_heading(level=level)
        r = par.add_run(text)
        set_font(r, SERIF, 16 if level == 1 else 14, bold=True)
        r.font.color.rgb = RGBColor(0, 0, 0)
        self.headings.append((level, text))
        return par

    def bullets(self, items, style="List Bullet"):
        for it in items:
            par = self.d.add_paragraph(style=style)
            par.paragraph_format.first_line_indent = Cm(0)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self._runs(par, it)

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            par = self.d.add_paragraph()
            par.paragraph_format.left_indent = Cm(0)
            self._runs(par, f"{i}) {it}")

    def page_break(self):
        par = self.d.add_paragraph()
        par.paragraph_format.first_line_indent = Cm(0)
        par.add_run().add_break(WD_BREAK.PAGE)

    def code(self, text, size=10.5):
        lines = text.split("\n")
        for i, line in enumerate(lines):
            par = self.d.add_paragraph()
            pf = par.paragraph_format
            pf.first_line_indent = Cm(0); pf.left_indent = Cm(0.3)
            pf.line_spacing = 1.0
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.keep_with_next = i < len(lines) - 1 and len(lines) < 30
            if i == 0: pf.space_before = Pt(4)
            if i == len(lines) - 1: pf.space_after = Pt(8)
            para_shade(par, "F2F2F2")
            r = par.add_run(line if line else " ")
            set_font(r, MONO, size)

    def table(self, header, rows, widths, caption=None, size=12, header_fill="D9E2F3", bold_first_col=False):
        if caption:
            self.tab += 1
            cp = self.p(f"Таблица {self.tab} — {caption}", indent=False, align="left", keep=True, space_after=3)
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        def fill(cell, text, bold=False):
            cell.text = ""
            lines = str(text).split("\n")
            for j, line in enumerate(lines):
                par = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                pf = par.paragraph_format
                pf.first_line_indent = Cm(0); pf.line_spacing = 1.0
                pf.space_after = Pt(1); pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._runs(par, line, bold=bold, size=size)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        hdr = t.rows[0]
        # repeat header row
        trPr = hdr._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
        for i, htxt in enumerate(header):
            fill(hdr.cells[i], htxt, bold=True); shade(hdr.cells[i], header_fill)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                fill(cells[i], v, bold=(bold_first_col and i == 0))
        for row in t.rows:
            # don't split rows across pages
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        tblPr = t._tbl.tblPr
        tw_ = tblPr.find(qn("w:tblW"))
        if tw_ is None:
            tw_ = OxmlElement("w:tblW"); tblPr.append(tw_)
        tw_.set(qn("w:w"), str(int(sum(widths) * 567))); tw_.set(qn("w:type"), "dxa")
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        sp = self.d.add_paragraph(); sp.paragraph_format.first_line_indent = Cm(0)
        sp.paragraph_format.line_spacing = 1.0
        return t

    def image(self, path, width_cm, caption):
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.space_before = Pt(6)
        par.add_run().add_picture(path, width=Cm(width_cm))
        self.fig += 1
        c = self.p(f"Рисунок {self.fig} — {caption}", indent=False, align="center", space_after=8)
        c.paragraph_format.line_spacing = 1.15
        return self.fig

    def landscape(self, on=True):
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        if on:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
            sec.left_margin, sec.right_margin = Cm(2.0), Cm(2.0)
            sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(1.5)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
            sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
            sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        sec.different_first_page_header_footer = False
        return sec

    # ------------------------------------------------------------ title page
    def title_page(self, c, work_title, subtitle):
        def line(text, bold=False, size=14, after=0, align="center"):
            par = self.p(text, bold=bold, align=align, indent=False, size=size, space_after=after)
            par.paragraph_format.line_spacing = 1.15
            return par
        line(c.INSTITUTE.upper(), bold=True)
        line(c.DEPARTMENT, after=60)
        line(f"Дисциплина: «{c.DISCIPLINE}»", after=48)
        line("ПРАКТИЧЕСКОЕ ЗАДАНИЕ", bold=True, size=16, after=6)
        line(work_title, bold=True, size=16, after=24)
        line("Тема проекта № 12:", after=4)
        line(f"«{c.TOPIC}»", bold=True, after=4)
        line(f"(рабочее название — {c.PROJECT})", after=4)
        if subtitle:
            line(subtitle, after=0)
        for _ in range(5):
            line("")
        t = self.d.add_table(rows=3, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.RIGHT
        data = [("Выполнили:", f"студенты группы {c.GROUP}\n" + "\n".join(c.STUDENTS)),
                ("Преподаватель:", c.TEACHER),
                ("Семестр:", "1 семестр, 2026–2027 уч. г.")]
        for i, (k, v) in enumerate(data):
            for j, txt in enumerate((k, v)):
                cell = t.rows[i].cells[j]
                cell.text = ""
                for n, ln in enumerate(txt.split("\n")):
                    par = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
                    par.paragraph_format.first_line_indent = Cm(0)
                    par.paragraph_format.line_spacing = 1.15
                    par.paragraph_format.space_after = Pt(6 if n == len(txt.split("\n")) - 1 else 0)
                    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    r = par.add_run(ln); set_font(r, SERIF, 14, bold=(j == 0))
            t.rows[i].cells[0].width = Cm(4.0); t.rows[i].cells[1].width = Cm(10.0)
        t.columns[0].width = Cm(4.0); t.columns[1].width = Cm(10.0)
        # spacer to push city/year to the bottom
        for _ in range(4):
            line("")
        par = line(c.CITY_YEAR)
        self.page_break()

    # ------------------------------------------------------------ TOC (filled in second pass)
    def toc(self):
        self.p("СОДЕРЖАНИЕ", bold=True, align="center", indent=False, space_after=12)
        self._toc_anchor = self.d.add_paragraph()
        self._toc_anchor.paragraph_format.first_line_indent = Cm(0)
        self.page_break()

    def fill_toc(self, pages):
        anchor = self._toc_anchor
        for level, text in self.headings:
            par = OxmlElement("w:p")
            anchor._p.addprevious(par)
            from docx.text.paragraph import Paragraph
            P = Paragraph(par, anchor._parent)
            pf = P.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.left_indent = Cm(0 if level == 1 else 0.8)
            pf.line_spacing = 1.25
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            sec = self.d.sections[0]
            width = sec.page_width - sec.left_margin - sec.right_margin
            pf.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            r = P.add_run(f"{text}\t{pages.get(text, '')}")
            set_font(r, SERIF, 14, bold=(level == 1))

    def save(self, path):
        self.d.save(path)


def to_pdf(docx_path, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")


def heading_pages(pdf, headings, start_page=3):
    out = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    pages = out.split("\f")
    norm = lambda s: re.sub(r"\s+", " ", s).strip()
    res = {}
    cursor = start_page - 1
    for level, text in headings:
        key = norm(text)[:45]
        for i in range(cursor, len(pages)):
            if key in norm(pages[i]):
                res[text] = i + 1
                cursor = i
                break
    return res
